import openpyxl
import json
import re
from datetime import datetime

SRC = "/Users/yapfongkiat/Downloads/S2 Student Data (2).xlsx"

HOSTELS = [
    {"code": "ATR", "name": "Atria", "sheet": "Atria", "address": "JALAN NOVA U5/71A, SUBANG BESTARI, 40150 SHAH ALAM, SELANGOR D. E., MALAYSIA.", "electricity_rate": 0.685},
    {"code": "DAM", "name": "Damai", "sheet": "Damai", "address": "JALAN TASIK RAJA LUMU U4/17, TAMAN SUBANG DELIMA, 40150 SHAH ALAM, SELANGOR D. E., MALAYSIA.", "electricity_rate": 0.685},
    {"code": "NDY", "name": "Nadayu 801", "sheet": "Nadayu 801", "address": "1, JALAN ZUHRAH U5/163, SUBANG MURNI, SEKSYEN U5, 40150 SHAH ALAM, SELANGOR D.E., MALAYSIA.", "electricity_rate": 0.751},
    {"code": "SHP", "name": "Shop Hostel", "sheet": "Shop Hostel", "address": "16, JALAN NOVA U5/N, SUBANG BESTARI, 40150 SHAH ALAM, SELANGOR D. E., MALAYSIA.", "electricity_rate": 0.685},
    {"code": "SR", "name": "Subang Residences", "sheet": "Subang Residence", "address": "JALAN DALWU U5/98, 40150 SHAH ALAM, SELANGOR D. E., MALAYSIA.", "electricity_rate": 0.685},
]

# Shop Hostel gender split (only hostel with a real gender split, per Addresses sheet)
SHOP_GENDER = {"16-2": "female", "16-3": "male"}


def split_unit_code(unit_str):
    """'D1-0113-A1' -> ('D1-0113', 'A', '1'); '1304-A1' -> ('1304','A','1')"""
    parts = unit_str.strip().split("-")
    last = parts[-1]
    m = re.match(r"^([A-Za-z]+)(\d+)$", last)
    if not m:
        return None
    room_label, bed_label = m.group(1), m.group(2)
    unit_code = "-".join(parts[:-1])
    return unit_code, room_label, bed_label


def fmt_date(val):
    if val is None:
        return None
    if isinstance(val, datetime):
        return val.strftime("%Y-%m-%d")
    s = str(val).strip()
    return s or None


def parse_tenant_name(raw):
    """'Hua Lai Yee (B2500480)' -> ('Hua Lai Yee', 'B2500480')"""
    if not raw:
        return None, None
    raw = str(raw).strip()
    if raw.startswith("(") and raw.endswith(")"):
        # e.g. "(Currently Storeroom)" - not a real tenant
        return None, None
    m = re.match(r"^(.*?)\s*\(([^)]*)\)\s*$", raw)
    if m:
        return m.group(1).strip(), m.group(2).strip()
    return raw, None


def num(v):
    if v is None or v == "":
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


wb = openpyxl.load_workbook(SRC, data_only=True)

result_hostels = []
all_beds = []
all_students = []  # dedup by (name, ic) later
skipped_rows = []

for h in HOSTELS:
    ws = wb[h["sheet"]]
    header = [c for c in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
    col = {name: idx for idx, name in enumerate(header) if name}

    units_map = {}  # unit_code -> {rooms: {room_label: {beds: {...}}}}

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=True):
        unit_raw = row[col["UNIT"]]
        if not unit_raw:
            continue
        parsed = split_unit_code(str(unit_raw))
        if not parsed:
            skipped_rows.append({"hostel": h["code"], "unit_raw": unit_raw, "reason": "unparsable code"})
            continue
        unit_code, room_label, bed_label = parsed

        tenant_raw = row[col.get("TENANT NAME")] if "TENANT NAME" in col else None
        tenant_name, student_code = parse_tenant_name(tenant_raw)
        is_storeroom = tenant_raw and str(tenant_raw).strip().lower().startswith("(currently")

        bed_entry = {
            "legacy_code": str(unit_raw).strip(),
            "monthly_rental": num(row[col["MONTHLY RENTAL"]]) if "MONTHLY RENTAL" in col else None,
            "access_card_deposit": num(row[col["ACCESS CARD DEPOSIT"]]) if "ACCESS CARD DEPOSIT" in col else None,
            "special_use": "storeroom" if is_storeroom else None,
        }

        assignment = None
        if tenant_name:
            ci_date = fmt_date(row[col["CI DATE"]]) if "CI DATE" in col else None
            co_date = fmt_date(row[col["CO DATE"]]) if "CO DATE" in col else None
            assignment = {
                "student": {
                    "source_code": student_code or "",
                    "full_name": tenant_name,
                    "ic_passport": row[col.get("IC / PASSPORT NO")] if "IC / PASSPORT NO" in col else None,
                    "contact": row[col.get("CONTACT")] if "CONTACT" in col else None,
                    "nationality": row[col.get("NATIONALITY")] if "NATIONALITY" in col else None,
                    "hometown": row[col.get("HOMETOWN")] if "HOMETOWN" in col else None,
                    "course": row[col.get("COURSE")] if "COURSE" in col else None,
                },
                "monthly_rental": num(row[col["MONTHLY RENTAL"]]) if "MONTHLY RENTAL" in col else None,
                "security_deposit": num(row[col["SECURITY DEPOSIT"]]) if "SECURITY DEPOSIT" in col else None,
                "access_card_deposit": num(row[col["ACCESS CARD DEPOSIT"]]) if "ACCESS CARD DEPOSIT" in col else None,
                "salesperson": (row[col.get("SALESPERSON")] or "").strip() if "SALESPERSON" in col and row[col.get("SALESPERSON")] else "",
                "check_in_date": ci_date,
                "agreement_start_date": fmt_date(row[col["LATEST TA START DATE"]]) if "LATEST TA START DATE" in col else None,
                "agreement_end_date": fmt_date(row[col["TA END ON"]]) if "TA END ON" in col else None,
                "agreement_duration": str(row[col["LATEST TA DURATION"]]) if "LATEST TA DURATION" in col and row[col["LATEST TA DURATION"]] else "",
                "check_out_date": co_date,
                "check_in_meter": num(row[col["CI METER"]]) if "CI METER" in col else None,
                "remarks": str(row[col["REMARK(S)"]]) if "REMARK(S)" in col and row[col["REMARK(S)"]] else "",
                "status": "ended" if co_date else "active",
            }

        units_map.setdefault(unit_code, {})
        units_map[unit_code].setdefault(room_label, {})
        units_map[unit_code][room_label][bed_label] = {"bed": bed_entry, "assignment": assignment}

    # Build unit/room/bed structure + gender rule
    unit_list = []
    for unit_code, rooms in units_map.items():
        if h["code"] == "SHP":
            gender = SHOP_GENDER.get(unit_code, "mixed")
        else:
            gender = "mixed"
        room_list = []
        for room_label, beds in rooms.items():
            bed_list = []
            for bed_label, entry in beds.items():
                bed_list.append({"bed_label": bed_label, **entry["bed"], "assignment": entry["assignment"]})
            room_list.append({"room_label": room_label, "beds": sorted(bed_list, key=lambda b: b["bed_label"])})
        unit_list.append({"unit_code": unit_code, "gender": gender, "rooms": sorted(room_list, key=lambda r: r["room_label"])})

    result_hostels.append({
        "code": h["code"],
        "name": h["name"],
        "address": h["address"],
        "electricity_rate": h["electricity_rate"],
        "units": sorted(unit_list, key=lambda u: u["unit_code"]),
    })

# Stats
total_beds = 0
total_active = 0
total_ended = 0
total_vacant = 0
total_storeroom = 0
for h in result_hostels:
    for u in h["units"]:
        for r in u["rooms"]:
            for b in r["beds"]:
                total_beds += 1
                if b["special_use"] == "storeroom":
                    total_storeroom += 1
                elif b["assignment"] is None:
                    total_vacant += 1
                elif b["assignment"]["status"] == "active":
                    total_active += 1
                else:
                    total_ended += 1

print(json.dumps({
    "total_beds": total_beds,
    "total_active_tenants": total_active,
    "total_ended_tenants": total_ended,
    "total_vacant": total_vacant,
    "total_storeroom": total_storeroom,
    "skipped_rows": len(skipped_rows),
}, indent=2))

if skipped_rows:
    print("SKIPPED SAMPLE:", skipped_rows[:10])

with open("/Users/yapfongkiat/create-a-system-database-for-my/.claude-scratch/import-data.json", "w") as f:
    json.dump(result_hostels, f, indent=2, default=str)
print("wrote import-data.json")
