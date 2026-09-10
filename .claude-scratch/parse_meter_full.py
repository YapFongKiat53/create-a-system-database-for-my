import openpyxl
import json
from datetime import datetime

SRC = "/Users/yapfongkiat/Downloads/Meter Reading 2026.xlsx"
wb = openpyxl.load_workbook(SRC, data_only=True)


def parse_sheet(name):
    ws = wb[name]
    current_dates = None  # list of (col_index, date_str)
    readings = []  # (code, date_str, value)
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=True):
        first = row[0]
        if first == "Date":
            current_dates = []
            for i, v in enumerate(row[1:], start=1):
                if isinstance(v, datetime):
                    current_dates.append((i, v.strftime("%Y-%m-%d")))
            continue
        if not first or not isinstance(first, str) or first.strip() in ("", "-"):
            continue
        code = first.strip()
        if current_dates is None:
            continue
        for col_i, date_str in current_dates:
            val = row[col_i] if col_i < len(row) else None
            if val is not None and isinstance(val, (int, float)):
                readings.append({"code": code, "date": date_str, "value": val})
    return readings


all_stats = {}
all_readings = {}
for sheet in ["ATRIA", "DAMAI", "DAMAI 1", "NADAYU", "SR, SHOPLOT"]:
    r = parse_sheet(sheet)
    dates = sorted(set(x["date"] for x in r))
    codes = sorted(set(x["code"] for x in r))
    all_stats[sheet] = {
        "total_readings": len(r),
        "unique_codes": len(codes),
        "date_range": [dates[0], dates[-1]] if dates else None,
        "sample_codes": codes[:10],
    }
    all_readings[sheet] = r

print(json.dumps(all_stats, indent=2))
with open("/Users/yapfongkiat/create-a-system-database-for-my/.claude-scratch/meter-readings-raw.json", "w") as f:
    json.dump(all_readings, f)
