import openpyxl
import json
import re
from datetime import datetime

SRC = "/Users/yapfongkiat/Downloads/S2 Parking.xlsx"

HOSTEL_SHEETS = {
    "Atria": "ATR",
    "Damai": "DAM",
    "Nadayu": "NDY",
}

wb = openpyxl.load_workbook(SRC, data_only=True)

out = []
for sheet_name, hostel_code in HOSTEL_SHEETS.items():
    ws = wb[sheet_name]
    header = [c for c in next(ws.iter_rows(min_row=3, max_row=3, values_only=True))]
    col = {n: i for i, n in enumerate(header) if n}

    for row in ws.iter_rows(min_row=4, max_row=ws.max_row, values_only=True):
        lot_no = row[col["P. LOT NO"]] if "P. LOT NO" in col else None
        if lot_no is None:
            continue
        name_raw = row[col["NAME"]] if "NAME" in col else None
        is_placeholder = name_raw and str(name_raw).strip().lower().startswith("(taken by")
        rental = {
            "hostel_code": hostel_code,
            "lot_number": str(lot_no).strip(),
            "lot_unit": str(row[col["P. LOT UNIT"]]).strip() if row[col.get("P. LOT UNIT", -1)] else "",
            "access_card_no": (
                row[col["ACCESS CARD NO"]] if "ACCESS CARD NO" in col
                else row[col["RFID NO"]] if "RFID NO" in col else None
            ),
            "house_unit": str(row[col["HOUSE UNIT"]]).strip() if row[col.get("HOUSE UNIT", -1)] not in (None, "-") else None,
            "tenant_name": None if (not name_raw or is_placeholder) else str(name_raw).strip(),
            "contact": row[col.get("CONTACT")] if "CONTACT" in col else None,
            "monthly_rental": row[col["RENTAL"]] if "RENTAL" in col and row[col["RENTAL"]] is not None else None,
            "deposit": row[col["DEPOSIT"]] if "DEPOSIT" in col and row[col["DEPOSIT"]] is not None else None,
            "car_plate": row[col.get("CAR PLAT NO")] if "CAR PLAT NO" in col else None,
            "car_model": row[col.get("CAR MODEL")] if "CAR MODEL" in col else None,
            "remark": row[col.get("REMARK")] if "REMARK" in col else None,
            "returning_date": row[col.get("RETURNING DATE")] if "RETURNING DATE" in col else None,
        }
        out.append(rental)

total = len(out)
rented = sum(1 for r in out if r["tenant_name"])
print(json.dumps({"total_lots": total, "rented": rented}, indent=2))

with open("/Users/yapfongkiat/create-a-system-database-for-my/.claude-scratch/parking-data.json", "w") as f:
    json.dump(out, f, indent=2, default=str)
print("wrote parking-data.json")
